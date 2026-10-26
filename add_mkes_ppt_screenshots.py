"""
add_mkes_ppt_screenshots.py

Automated PowerPoint Screenshot Integration for MKES CARE+ Presentation.
Embeds 11 high-fidelity mobile application screenshots into slides 1, 12, 13, and 14
of the official presentation while strictly preserving all original slides, formatting,
fonts, titles, and non-target shapes.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw


def round_corners(image: Image.Image, radius: int) -> Image.Image:
    """Apply smooth rounded corners to an image with an RGBA alpha mask."""
    mask = Image.new('L', image.size, 0)
    draw = ImageDraw.Draw(mask)
    draw.rounded_rectangle([(0, 0), image.size], radius=radius, fill=255)
    image_rgba = image.convert('RGBA')
    image_rgba.putalpha(mask)
    return image_rgba


def main():
    print("=" * 70)
    print("MKES CARE+ PowerPoint Screenshot Integration")
    print("=" * 70)

    # 1. Path definitions
    workspace_dir = r"d:\mkes-lovable"
    downloads_dir = r"C:\Users\Tatva jain\Downloads"

    # Input presentation path
    input_pptx = os.path.join(downloads_dir, "MKES CARE+ Front Page (1).pptx")
    if not os.path.exists(input_pptx):
        # Fallback check in workspace
        alt_input = os.path.join(workspace_dir, "MKES CARE+ Front Page (1).pptx")
        if os.path.exists(alt_input):
            input_pptx = alt_input
        else:
            raise FileNotFoundError(f"Original presentation not found at: {input_pptx}")

    # Output presentation path
    output_pptx = os.path.join(workspace_dir, "MKES_Care_Plus_Presentation_Screenshots_Added.pptx")

    # Temp directory for processed transparent-corner screenshot assets
    temp_dir = os.path.join(workspace_dir, "temp_ppt_assets")
    os.makedirs(temp_dir, exist_ok=True)

    # 2. Screenshot mappings & existence checks
    screenshot_sources = {
        "s1_dashboard": os.path.join(workspace_dir, "presentation_screenshots", "final_selected", "Admin", "admin_dashboard.png"),
        "s12_auth": os.path.join(workspace_dir, "image_e96700.png"),
        "s12_dashboard": os.path.join(workspace_dir, "presentation_screenshots", "final_selected", "Admin", "admin_dashboard.png"),
        "s12_queues": os.path.join(workspace_dir, "presentation_screenshots", "final_selected", "Staff", "staff_dashboard.png"),
        "s12_staff": os.path.join(workspace_dir, "presentation_screenshots", "final_selected", "Admin", "admin_staff_management.png"),
        "s12_inventory": os.path.join(workspace_dir, "presentation_screenshots", "final_selected", "Admin", "admin_inventory.png"),
        "s12_billing": os.path.join(workspace_dir, "presentation_screenshots", "final_selected", "Admin", "admin_reports_analytics.png"),
        "s13_surgery": os.path.join(workspace_dir, "presentation_screenshots", "final_selected", "Admin", "admin_surgery_management.png"),
        "s13_patient": os.path.join(workspace_dir, "presentation_screenshots", "final_selected", "Staff", "staff_patient_management.png"),
        "s14_shifts": os.path.join(workspace_dir, "presentation_screenshots", "final_selected", "Admin", "admin_staff_shifts_and_scheduling.png"),
        "s14_tasks": os.path.join(workspace_dir, "presentation_screenshots", "final_selected", "Staff", "staff_assigned_surgery_tasks.png"),
        "s14_usage": os.path.join(workspace_dir, "presentation_screenshots", "final_selected", "Staff", "staff_inventory_stock_usage.png"),
    }

    print("[Step 1] Verifying input files...")
    print(f"  Original PPTX: {input_pptx}")
    for key, path in screenshot_sources.items():
        if not os.path.exists(path):
            raise FileNotFoundError(f"Missing required screenshot '{key}': {path}")
        print(f"  [OK] {key:15}: {path}")

    # Check original file size and modification time to ensure preservation
    orig_stat = os.stat(input_pptx)
    orig_mtime = orig_stat.st_mtime
    orig_size = orig_stat.st_size
    print(f"  Original PPTX verified (Size: {orig_size} bytes, MTime: {orig_mtime})")

    # Load presentation
    prs = Presentation(input_pptx)
    print(f"\nLoaded presentation with {len(prs.slides)} slides.")

    # =========================================================================
    # SLIDE 1: Title Slide Phone Mockup
    # =========================================================================
    print("\n[Step 2] Processing Slide 1...")
    s1 = prs.slides[0]

    # Remove placeholder shapes (Picture 16 and Picture 18)
    removed_s1 = []
    for sh in list(s1.shapes):
        if sh.name in ['Picture 16', 'Picture 18']:
            sp = sh._element
            sp.getparent().remove(sp)
            removed_s1.append(sh.name)
    print(f"  Removed placeholders: {removed_s1}")

    # Prepare admin_dashboard.png with matching aspect ratio crop and rounded corners
    im1 = Image.open(screenshot_sources["s1_dashboard"])
    crop_h1 = int(im1.width * (5276850 / 3124200))  # 1317 px for 1.689 aspect ratio
    im1_cropped = im1.crop((0, 0, im1.width, crop_h1))
    im1_rounded = round_corners(im1_cropped, radius=55)
    im1_path = os.path.join(temp_dir, "slide_01_dashboard.png")
    im1_rounded.save(im1_path)

    # Insert into phone screen area
    pic1 = s1.shapes.add_picture(im1_path, 7972425, 1581150, 3124200, 5276850)

    # Reorder in XML shape tree immediately after Rounded Rectangle 6
    # Keeps top notch (Picture 15) and bottom pill (Rounded Rectangle 10) on top
    rr6_elem = None
    for sh in s1.shapes:
        if sh.name == 'Rounded Rectangle 6':
            rr6_elem = sh._element
            break
    if rr6_elem is not None:
        spTree = s1.shapes._spTree
        spTree.remove(pic1._element)
        rr6_idx = list(spTree).index(rr6_elem)
        spTree.insert(rr6_idx + 1, pic1._element)
    print("  Inserted admin_dashboard.png into Slide 1 phone mockup.")

    # =========================================================================
    # SLIDE 12: Core Modules (6 Phone Mockups)
    # =========================================================================
    print("\n[Step 3] Processing Slide 12...")
    s12 = prs.slides[11]

    # Remove placeholder icons and slot text boxes
    to_remove_12 = [
        'Picture 22', 'Picture 23', 'Picture 24', 'Picture 25', 'Picture 26', 'Picture 27',
        'TextBox 44', 'TextBox 45', 'TextBox 46', 'TextBox 47', 'TextBox 48', 'TextBox 49'
    ]
    removed_s12 = []
    for sh in list(s12.shapes):
        if sh.name in to_remove_12:
            sp = sh._element
            sp.getparent().remove(sp)
            removed_s12.append(sh.name)
    print(f"  Removed placeholders: {removed_s12}")

    s12_config = [
        ("Auth Screen", screenshot_sources["s12_auth"], 858738, 2162175, 1219200, 2266950, 'Rounded Rectangle 15'),
        ("Dashboard", screenshot_sources["s12_dashboard"], 2709713, 2147887, 1219200, 2266950, 'Rounded Rectangle 16'),
        ("Queue UI", screenshot_sources["s12_queues"], 4560689, 2247900, 1219200, 2266950, 'Rounded Rectangle 17'),
        ("Staff UI", screenshot_sources["s12_staff"], 6411664, 2247900, 1219200, 2266950, 'Rounded Rectangle 18'),
        ("Inventory", screenshot_sources["s12_inventory"], 8262639, 2247900, 1219200, 2266950, 'Rounded Rectangle 19'),
        ("Billing UI", screenshot_sources["s12_billing"], 10113615, 2162175, 1219200, 2266950, 'Rounded Rectangle 20'),
    ]

    for idx, (label, img_path, left, top, w, h, screen_name) in enumerate(s12_config):
        im = Image.open(img_path)
        crop_h = int(im.width * (h / w))
        im_cropped = im.crop((0, 0, im.width, crop_h))
        im_rounded = round_corners(im_cropped, radius=35)
        save_p = os.path.join(temp_dir, f"slide_12_slot_{idx+1}.png")
        im_rounded.save(save_p)

        pic = s12.shapes.add_picture(save_p, left, top, w, h)

        # Place immediately after the inner screen shape so outer bezels stay on top
        screen_elem = None
        for sh in s12.shapes:
            if sh.name == screen_name:
                screen_elem = sh._element
                break
        if screen_elem is not None:
            spTree = s12.shapes._spTree
            spTree.remove(pic._element)
            screen_idx = list(spTree).index(screen_elem)
            spTree.insert(screen_idx + 1, pic._element)
        print(f"  Slot {idx+1} ({label}): inserted {os.path.basename(img_path)}")

    # =========================================================================
    # SLIDE 13: Clinical Modules (Eye & Dental Care)
    # =========================================================================
    print("\n[Step 4] Processing Slide 13...")
    s13 = prs.slides[12]

    # Remove only Picture 8 (the stock illustration)
    removed_s13 = []
    for sh in list(s13.shapes):
        if sh.name == 'Picture 8':
            sp = sh._element
            sp.getparent().remove(sp)
            removed_s13.append(sh.name)
    print(f"  Removed stock graphic: {removed_s13}")

    s13_phones = [
        ("Eye Surgery", screenshot_sources["s13_surgery"], 4914000, 2480000, 1102000, 2280000, 4959000, 2525000, 1012000, 2190000),
        ("Patient Directory", screenshot_sources["s13_patient"], 6176000, 2480000, 1102000, 2280000, 6221000, 2525000, 1012000, 2190000),
    ]

    for idx, (label, img_path, b_left, b_top, b_w, b_h, s_left, s_top, s_w, s_h) in enumerate(s13_phones):
        # Outer bezel shape
        bezel = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b_left, b_top, b_w, b_h)
        bezel.fill.solid()
        bezel.fill.fore_color.rgb = RGBColor(17, 17, 17)
        bezel.line.color.rgb = RGBColor(40, 40, 40)
        bezel.line.width = Pt(1)

        # Screenshot with natural aspect ratio (2.164:1) and rounded corners
        im = Image.open(img_path)
        im_rounded = round_corners(im, radius=32)
        save_p = os.path.join(temp_dir, f"slide_13_phone_{idx+1}.png")
        im_rounded.save(save_p)
        pic = s13.shapes.add_picture(save_p, s_left, s_top, s_w, s_h)

        # Notch pill
        notch_w = 260000
        notch_h = 35000
        notch_left = b_left + (b_w - notch_w) // 2
        notch_top = s_top + 10000
        notch = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, notch_left, notch_top, notch_w, notch_h)
        notch.fill.solid()
        notch.fill.fore_color.rgb = RGBColor(17, 17, 17)
        notch.line.fill.background()
        print(f"  Inserted phone mockup {idx+1} ({label}): {os.path.basename(img_path)}")

    # =========================================================================
    # SLIDE 14: Workforce & Supply Chain Control
    # =========================================================================
    print("\n[Step 5] Processing Slide 14...")
    s14 = prs.slides[13]

    # Remove only Picture 12 and Picture 13 (stock spreadsheet graphic)
    removed_s14 = []
    for sh in list(s14.shapes):
        if sh.name in ['Picture 12', 'Picture 13']:
            sp = sh._element
            sp.getparent().remove(sp)
            removed_s14.append(sh.name)
    print(f"  Removed stock graphic: {removed_s14}")

    s14_phones = [
        ("Shift Scheduling", screenshot_sources["s14_shifts"], 7787775, 2310000, 981000, 2030000, 7827775, 2350000, 901000, 1950000),
        ("Surgery Tasks", screenshot_sources["s14_tasks"], 8948775, 2310000, 981000, 2030000, 8988775, 2350000, 901000, 1950000),
        ("Stock Usage", screenshot_sources["s14_usage"], 10109775, 2310000, 981000, 2030000, 10149775, 2350000, 901000, 1950000),
    ]

    for idx, (label, img_path, b_left, b_top, b_w, b_h, s_left, s_top, s_w, s_h) in enumerate(s14_phones):
        # Outer bezel shape
        bezel = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b_left, b_top, b_w, b_h)
        bezel.fill.solid()
        bezel.fill.fore_color.rgb = RGBColor(17, 17, 17)
        bezel.line.color.rgb = RGBColor(40, 40, 40)
        bezel.line.width = Pt(1)

        # Screenshot with natural aspect ratio (2.164:1) and rounded corners
        im = Image.open(img_path)
        im_rounded = round_corners(im, radius=28)
        save_p = os.path.join(temp_dir, f"slide_14_phone_{idx+1}.png")
        im_rounded.save(save_p)
        pic = s14.shapes.add_picture(save_p, s_left, s_top, s_w, s_h)

        # Notch pill
        notch_w = 230000
        notch_h = 32000
        notch_left = b_left + (b_w - notch_w) // 2
        notch_top = s_top + 8000
        notch = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, notch_left, notch_top, notch_w, notch_h)
        notch.fill.solid()
        notch.fill.fore_color.rgb = RGBColor(17, 17, 17)
        notch.line.fill.background()
        print(f"  Inserted phone mockup {idx+1} ({label}): {os.path.basename(img_path)}")

    # =========================================================================
    # SAVE AND PRESERVATION VERIFICATION
    # =========================================================================
    print("\n[Step 6] Saving presentation...")
    prs.save(output_pptx)
    print(f"  Successfully saved to: {output_pptx}")
    print(f"  Output size: {os.path.getsize(output_pptx)} bytes")

    # Strict check: ensure original file was NOT modified
    after_stat = os.stat(input_pptx)
    if after_stat.st_mtime != orig_mtime or after_stat.st_size != orig_size:
        raise RuntimeError("CRITICAL ERROR: Original presentation was modified!")
    print("  [PRESERVATION CHECK PASSED] Original presentation remains completely untouched.")
    print("=" * 70)


if __name__ == "__main__":
    main()
